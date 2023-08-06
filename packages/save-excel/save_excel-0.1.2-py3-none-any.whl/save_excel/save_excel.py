from .config import Functionalities, Properties, TableStyles
from .dictionary import factory_dict
from .utils import set_value, validate_path, get_column_names
from .layout import ExcelLayout

from os import makedirs, path
from pandas import ExcelWriter
from shutil import copy
from copy import deepcopy
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation


class ExcelSaver:
    
    def __init__(self, obj_to_save=None, bookname='Book1', sheetname='Sheet1'):
        
        self.__xl_dict = {}
        self.__metadata = {}
        self.__styles = TableStyles()
        self.__properties = Properties()
        self.__functions = Functionalities()
        
        if obj_to_save is not None:
            self.add_tables(obj=obj_to_save, bookname=bookname, sheetname=sheetname)
    
    def add_tables(self, obj, bookname='Book1', sheetname='Sheet1', position='right', level=-1):
        
        builded_dict = factory_dict(obj, bookname=bookname, sheetname=sheetname).dict
        
        for bookname in builded_dict:
            if bookname in self.__xl_dict:
                for sheetname in builded_dict[bookname]:
                    if sheetname in self.__xl_dict[bookname]:
                        for n, df_iter in enumerate(builded_dict[bookname][sheetname]):
                            if position=='right' and n==0:
                                if isinstance(df_iter, dict) and isinstance(self.__xl_dict[bookname][sheetname][level], dict):
                                    self.__xl_dict[bookname][sheetname][level].update(df_iter)
                                elif isinstance(df_iter, list) and isinstance(self.__xl_dict[bookname][sheetname][level], list):
                                    self.__xl_dict[bookname][sheetname][level] += df_iter
                                else:
                                    raise ValueError("All tables in the same level must have title")
                            elif position=='bottom' or n>0:
                                if level<0:
                                    self.__xl_dict[bookname][sheetname].insert(len(self.__xl_dict[bookname][sheetname])-level+1, df_iter)
                                else:
                                    self.__xl_dict[bookname][sheetname].insert(level+n, df_iter)
                            else:
                                raise ValueError("Position options are: 'right', 'bottom'")
                    else:
                        self.__xl_dict[bookname][sheetname] = builded_dict[bookname][sheetname]
            else:
                self.__xl_dict[bookname] = builded_dict[bookname]
            
    def __generate_metadata(self):
        
        for bookname in self.__xl_dict:
            self.metadata[bookname] = {}
            for sheetname in self.__xl_dict[bookname]:

                self.metadata[bookname][sheetname] = ExcelLayout(data=self.__xl_dict[bookname][sheetname], 
                                                           index=self.properties.distribution.index, 
                                                           startcol=self.properties.distribution.startcol, 
                                                           startrow=self.properties.distribution.startrow, 
                                                           cols_between_tables=self.properties.distribution.cols_between_tables, 
                                                           rows_between_tables=self.properties.distribution.rows_between_tables
                                                          )
    
    def save(self, path='Report', template=None):
        
        files_saved = 0
        
        if path is not None:
            makedirs(path, exist_ok=True)
        
        path = validate_path(path)
        
        self.__generate_metadata()
        
        for bookname in self.metadata:
            
            files_saved += 1
            filename = f'{path}{bookname}.xlsx'
            
            if template is None:
                mode = 'w'
                if_sheet_exists = None
            else:
                mode = 'a'
                if_sheet_exists = 'overlay'
                copy(f'{path}{template}.xlsx', filename)
                
            with ExcelWriter(filename, engine='openpyxl', mode=mode, if_sheet_exists=if_sheet_exists) as writer:
                
                workbook = writer.book
                
                for sheetname in self.metadata[bookname]:
                    
                    data_layout = self.metadata[bookname][sheetname]
                    
                    for frame_layout in data_layout.data:
                        
                        startcol, startrow = frame_layout.header.coord

                        frame_layout.df.to_excel(writer, 
                                                 sheet_name=sheetname, 
                                                 index=self.properties.distribution.index, 
                                                 startcol=startcol, 
                                                 startrow=startrow)
                        # Styles
                        if self.styles is not None:

                            worksheet = workbook[sheetname]

                            # Title styles
                            if self.styles.title is not None:
                                for cell in frame_layout.title.cells:
                                    worksheet[cell].fill = self.styles.title.fill
                                    worksheet[cell].font = self.styles.title.font
                                    worksheet[cell].alignment = self.styles.title.alignment
                                    worksheet[cell].border = self.styles.title.border

                                    worksheet[cell] = frame_layout.title_name

                                for row in frame_layout.title.rows:
                                    worksheet.row_dimensions[row+1].height = self.styles.title.height

                                if frame_layout.title.range is not None:
                                    worksheet.merge_cells(frame_layout.title.range)
                                    
                            # Header styles
                            if self.styles.header is not None:
                                for cell in frame_layout.header.cells:
                                    worksheet[cell].fill = self.styles.header.fill
                                    worksheet[cell].font = self.styles.header.font
                                    worksheet[cell].alignment = self.styles.header.alignment
                                    worksheet[cell].border = self.styles.header.border

                                for row in frame_layout.header.rows:
                                    worksheet.row_dimensions[row+1].height = self.styles.header.height

                            # Cell styles
                            if self.styles.data is not None:
                                for cell in frame_layout.data.cells:
                                    worksheet[cell].fill = self.styles.data.fill
                                    worksheet[cell].font = self.styles.data.font
                                    worksheet[cell].alignment = self.styles.data.alignment
                                    worksheet[cell].border = self.styles.data.border

                                for row in frame_layout.data.rows:
                                    worksheet.row_dimensions[row+1].height = self.styles.data.height
                                    
                        # Functionalities
                        if self.functions is not None:

                            # Autofilter
                            if self.functions.autofilter==True:
                                worksheet.auto_filter.ref = frame_layout.header.range

                            # Conditional formatting
                            for col, rules_tmp in self.functions.conditional_format.items():

                                column_names = get_column_names(frame_layout.df)
                                
                                if col in column_names+frame_layout.df.columns.to_list():
                                    
                                    rules = deepcopy(rules_tmp)
                                    header_cell, data_cell_range = frame_layout.find_column(col)
                                    for rule in rules:
                                        if rule.type=='containsText' and rule.operator=='containsText':
                                            rule.formula[0] = rule.formula[0].format(data_cell_range)
                                        worksheet.conditional_formatting.add(data_cell_range, rule)
                            
                            # Comments
                            for col, comment in self.functions.comments.items():
                                if col in frame_layout.df.columns:
                                    header_cell, data_cell_range = frame_layout.find_column(col)
                                    worksheet[header_cell].comment = Comment(comment, "SMO", height=120, width=220)
                                    
                            # Data validation
                            for col, values_list in self.functions.data_validation.items():
                                if col in frame_layout.df.columns:
                                    header_cell, data_cell_range = frame_layout.find_column(col)
                                    value_str = '"{}"'.format(','.join(values_list))
                                    data_val = DataValidation(type="list", formula1=value_str)
                                    worksheet.add_data_validation(data_val)
                                    data_val.add(data_cell_range)
                                    
                    if self.functions is not None:

                        # Columns width
                        for col, width in data_layout.column_dimensions.items():
                            new_width = width + self.functions.columns_width.increase + 7
                            if self.functions.columns_width.min_limit is not None and new_width < self.functions.columns_width.min_limit:
                                new_width = self.functions.columns_width.min_limit
                            if self.functions.columns_width.max_limit is not None and new_width > self.functions.columns_width.max_limit:
                                new_width = self.functions.columns_width.max_limit
                            worksheet.column_dimensions[col].width = new_width

                    # Properties
                    if self.properties is not None:

                        # View properties
                        if self.properties.view is not None:
                            worksheet.sheet_view.view = self.properties.view.view
                            worksheet.sheet_view.showGridLines = self.properties.view.grid_lines
                            worksheet.freeze_panes = self.properties.view.freeze_panes
                        
                        # Printer properties
                        if self.properties.printer is not None:
                            worksheet.page_margins.left = self.properties.printer.margins.left
                            worksheet.page_margins.right = self.properties.printer.margins.right
                            worksheet.page_margins.top = self.properties.printer.margins.top
                            worksheet.page_margins.bottom = self.properties.printer.margins.bottom
                            worksheet.page_setup.orientation = self.properties.printer.orientation
                            worksheet.sheet_properties.pageSetUpPr.fitToPage = self.properties.printer.fit_to_page
                            worksheet.page_setup.fitToHeight = self.properties.printer.fit_to_height
                            worksheet.page_setup.fitToWidth = self.properties.printer.fit_to_width
                            worksheet.print_options.horizontalCentered = self.properties.printer.horizontal_centered
                            worksheet.print_options.verticalCentered = self.properties.printer.vertical_centered
        
        print(f'{files_saved} file(s) saved!')
        
    @property
    def metadata(self):
        return self.__metadata
    
    @property
    def styles(self):
        return self.__styles
    
    @styles.setter
    @set_value([None])
    def styles(self, value):
        self.__styles = value
    
    @property
    def properties(self):
        return self.__properties
    
    @properties.setter
    @set_value([None])
    def properties(self, value):
        self.__properties = value

    @property
    def functions(self):
        return self.__functions
    
    @functions.setter
    @set_value([None])
    def functions(self, value):
        self.__functions = value